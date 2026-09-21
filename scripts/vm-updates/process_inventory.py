#!/usr/bin/env python3
"""
신규입고 처리 스크립트
승인된 inventory_sessions을 읽어 바코드를 부여하고
공동판매 엑셀 파일 '입고' 시트에 신규 행을 추가합니다.

크론탭 등록:
  */2 * * * * /home/ubuntu/work_automation/run_inventory_process.sh >> /home/ubuntu/work_automation/logs/inventory_process.log 2>&1
"""
from __future__ import annotations

import logging
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Color, Font, PatternFill
from openpyxl.utils import column_index_from_string

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from jobs.dropbox_client import DropboxClient, DropboxClientError, DropboxConflictError

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s - %(message)s",
)
logger = logging.getLogger("inventory_process")

SUPABASE_URL  = os.getenv("SUPABASE_URL", "").strip().rstrip("/")
SUPABASE_KEY  = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "").strip()
DROPBOX_BASE  = os.getenv("DROPBOX_BASE_FOLDER", "/공동판매_auto").strip().rstrip("/")
WORK_DIR      = Path(os.getenv("INVENTORY_WORK_DIR", "/tmp/work_automation/inventory"))

# 입고 시트 열 정의 (실제 공동판매 파일 기준)
SHEET_NAME   = "입고"
COL_NUM      = "A"   # 넘버
COL_LOC      = "B"   # 위치
COL_DATE     = "C"   # 입고일
COL_NAME     = "D"   # 상품명 (=G&" "&H 수식 또는 직접 값)
COL_PRICE    = "E"   # 금액
COL_QTY      = "F"   # 수량
COL_BARCODE  = "G"   # 바코드
COL_CATEGORY = "H"   # 소분류
COL_CODE     = "I"   # 상품분류코드 (H열 기반 자동)
COL_NOTE     = "J"   # 비고
COL_SALE     = "K"   # 판매 (전매장매출 연동 수식)
COL_OUT      = "L"   # 출고
COL_STOCK    = "M"   # 재고 (=F-K)
COL_STOCK_AMT = "N"  # 재고금액 (=E*M)
COL_SOURCE   = "P"   # 출처 (거래처명, 행거의 source 필드)

# A~H열 수기입력 서식 (맑은 고딕 10pt, 가운데정렬) — 시트 컬럼 기본 스타일과 동일
CELL_FONT      = Font(name="맑은 고딕", size=10)
CELL_ALIGNMENT = Alignment(horizontal="center", vertical="center")

# B(위치)/C(입고일)열 배경색 — 수기입력 행과 동일한 채우기
FILL_LOC  = PatternFill(patternType="solid", fgColor=Color(rgb="FFCAEDFB"))
FILL_DATE = PatternFill(patternType="solid", fgColor=Color(theme=5, tint=0.7999816888943144))

# 온라인 입고라도 옷이 아닌 잡화는 H열 소분류에 "온" 접두사를 붙이지 않음
NON_CLOTHING_CATEGORIES = {"신발", "가방", "패션잡화", "넥타이", "벨트", "ETC", "패브릭", "모자"}

# 온라인 입고라도 특정 바코드 접두어 세션은 H열 소분류에 "온" 접두사를 붙이지 않음
NO_ON_BARCODE_PREFIXES = {"WWA"}

# H열 소분류 기록 시 카테고리명 대신 "ETC"로 통합 기록되는 비의류 카테고리 (모자/가방/패브릭은 제외, 실제 명칭 유지)
ETC_MAPPED_CATEGORIES = {"신발", "패션잡화", "넥타이", "벨트"}

# 매출 업데이트 시간대 충돌 방지 (update_move_locations.py 와 동일 규칙)
BLACKOUT_START = (0, 5)
BLACKOUT_END   = (0, 20)


# ── Supabase ─────────────────────────────────────────────────────────────────

def _sb_headers() -> dict:
    return {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
    }


def sb_get(table: str, params: dict) -> list[dict]:
    r = requests.get(
        f"{SUPABASE_URL}/rest/v1/{table}",
        headers=_sb_headers(), params=params, timeout=30,
    )
    r.raise_for_status()
    return r.json()


def sb_patch(table: str, match: dict, data: dict) -> None:
    params = {k: f"eq.{v}" for k, v in match.items()}
    r = requests.patch(
        f"{SUPABASE_URL}/rest/v1/{table}",
        headers=_sb_headers(), params=params, json=data, timeout=30,
    )
    r.raise_for_status()


# ── 유틸 ─────────────────────────────────────────────────────────────────────

def is_blackout() -> bool:
    now   = datetime.now()
    h, m  = now.hour, now.minute
    start = BLACKOUT_START[0] * 60 + BLACKOUT_START[1]
    end   = BLACKOUT_END[0]   * 60 + BLACKOUT_END[1]
    cur   = h * 60 + m
    return start <= cur <= end


def hanger_sort_key(hanger: dict):
    """hanger_number는 TEXT 컬럼이라 DB의 order="hanger_number.asc"는 사전식 정렬
    ("1","10","11","2",...)이 되어 버림 — 숫자면 숫자로, 아니면 문자열로 비교."""
    s = str(hanger.get("hanger_number") or "")
    return (0, int(s), "") if s.isdigit() else (1, 0, s)


def get_dropbox_folder(session_date: str) -> str:
    d = datetime.strptime(session_date, "%Y-%m-%d")
    return f"{DROPBOX_BASE}/{d.year}/{d.month}월"


def _copy_formula_row(formula: str, src_row: int, dst_row: int) -> str:
    """수식 내 비절대 행 번호(src_row)를 dst_row로 교체."""
    def replacer(m: re.Match) -> str:
        dollar_col = m.group(1)
        col        = m.group(2)
        dollar_row = m.group(3)
        row_num    = int(m.group(4))
        if dollar_row == "$":   # 절대 행 참조 — 변경 안 함
            return m.group(0)
        if row_num == src_row:
            return f"{dollar_col}{col}{dollar_row}{dst_row}"
        return m.group(0)
    return re.sub(r"(\$?)([A-Z]+)(\$?)(\d+)", replacer, formula)


def _read_formula(ws, row: int, col_letter: str) -> str | None:
    """셀에서 수식(= 시작) 문자열 반환. 없으면 None."""
    val = ws.cell(row=row, column=column_index_from_string(col_letter)).value
    if val and isinstance(val, str) and val.startswith("="):
        return val
    return None


def _find_last_data_row(ws) -> int:
    """G열(바코드)을 기준으로 마지막 데이터 행 반환. 없으면 1(헤더행)."""
    bc_col = column_index_from_string(COL_BARCODE)
    last   = 1
    for row in ws.iter_rows(min_row=2, max_col=bc_col, values_only=False):
        cell = row[bc_col - 1]
        if cell.value and str(cell.value).strip():
            last = cell.row
    return last


def _find_max_num(ws, last_data_row: int) -> int:
    """A열(넘버) 최댓값 반환."""
    num_col = column_index_from_string(COL_NUM)
    max_num = 0
    for row_idx in range(2, last_data_row + 1):
        val = ws.cell(row=row_idx, column=num_col).value
        try:
            if val is not None:
                max_num = max(max_num, int(val))
        except (ValueError, TypeError):
            pass
    return max_num


# ── 엑셀 쓰기 ────────────────────────────────────────────────────────────────

def write_inventory_rows(local_path: Path, items: list[dict]) -> int:
    """
    items: [{barcode, price, category, location, session_date, brand, product_name}]
    '입고' 시트에 신규 행 추가 후 저장. 작성 행 수 반환.
    """
    wb = load_workbook(local_path)   # data_only=False → 수식 읽기 가능
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(f"'{SHEET_NAME}' 시트 없음. 실제 시트: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    last_row = _find_last_data_row(ws)
    max_num  = _find_max_num(ws, last_row)

    # 마지막 행에서 수식 패턴 읽기
    formula_tmpl: dict[str, tuple[str, int]] = {}
    if last_row >= 2:
        for col in (COL_NAME, COL_CODE, COL_SALE, COL_OUT, COL_STOCK, COL_STOCK_AMT):
            f = _read_formula(ws, last_row, col)
            if f:
                formula_tmpl[col] = (f, last_row)

    write_row = last_row + 1
    written   = 0

    for item in items:
        row_num      = max_num + written + 1
        session_date = datetime.strptime(item["session_date"], "%Y-%m-%d")
        barcode      = item["barcode"]
        category       = item["category"]
        category_plain = item["category_plain"]
        brand          = item.get("brand") or ""

        # A: 넘버
        ws.cell(row=write_row, column=column_index_from_string(COL_NUM)).value = row_num

        # B: 위치
        ws.cell(row=write_row, column=column_index_from_string(COL_LOC)).value = item["location"]

        # C: 입고일 (datetime → 엑셀 날짜 타입)
        c_cell = ws.cell(row=write_row, column=column_index_from_string(COL_DATE))
        c_cell.value         = session_date
        c_cell.number_format = "YYYY.M.D"

        # D: 상품명 — 수식 패턴 있으면 복사, 없으면 직접 값
        if COL_NAME in formula_tmpl:
            tmpl, src = formula_tmpl[COL_NAME]
            ws.cell(row=write_row, column=column_index_from_string(COL_NAME)).value = \
                _copy_formula_row(tmpl, src, write_row)
        else:
            name = f"{barcode} {brand} {category_plain}" if brand else f"{barcode} {category_plain}"
            ws.cell(row=write_row, column=column_index_from_string(COL_NAME)).value = name

        # E: 금액 (천단위 콤마)
        price_cell = ws.cell(row=write_row, column=column_index_from_string(COL_PRICE))
        price_cell.value         = item["price"]
        price_cell.number_format = "#,##0"

        # F: 수량 (항상 1)
        ws.cell(row=write_row, column=column_index_from_string(COL_QTY)).value = 1

        # G: 바코드 (텍스트 형식)
        bc_cell = ws.cell(row=write_row, column=column_index_from_string(COL_BARCODE))
        bc_cell.value         = barcode
        bc_cell.number_format = "@"

        # H: 소분류
        ws.cell(row=write_row, column=column_index_from_string(COL_CATEGORY)).value = category

        # A~H: 폰트/정렬을 수기입력 서식(맑은 고딕 10pt, 가운데정렬)에 맞춤
        for col in range(1, 9):
            cell = ws.cell(row=write_row, column=col)
            cell.font      = CELL_FONT
            cell.alignment = CELL_ALIGNMENT

        # B/C: 수기입력 행과 동일한 배경색 유지
        ws.cell(row=write_row, column=column_index_from_string(COL_LOC)).fill  = FILL_LOC
        ws.cell(row=write_row, column=column_index_from_string(COL_DATE)).fill = FILL_DATE

        # I: 상품분류코드 — 수식 복사 우선
        if COL_CODE in formula_tmpl:
            tmpl, src = formula_tmpl[COL_CODE]
            ws.cell(row=write_row, column=column_index_from_string(COL_CODE)).value = \
                _copy_formula_row(tmpl, src, write_row)

        # J: 비고 (브랜드 — D열이 수식일 경우 J에 별도 기록)
        if brand and COL_NAME in formula_tmpl:
            ws.cell(row=write_row, column=column_index_from_string(COL_NOTE)).value = brand

        # K: 판매 (수식 복사)
        if COL_SALE in formula_tmpl:
            tmpl, src = formula_tmpl[COL_SALE]
            ws.cell(row=write_row, column=column_index_from_string(COL_SALE)).value = \
                _copy_formula_row(tmpl, src, write_row)

        # L: 출고 (수식 복사)
        if COL_OUT in formula_tmpl:
            tmpl, src = formula_tmpl[COL_OUT]
            ws.cell(row=write_row, column=column_index_from_string(COL_OUT)).value = \
                _copy_formula_row(tmpl, src, write_row)

        # M: 재고 — 수식 복사 또는 기본 =F-K
        if COL_STOCK in formula_tmpl:
            tmpl, src = formula_tmpl[COL_STOCK]
            ws.cell(row=write_row, column=column_index_from_string(COL_STOCK)).value = \
                _copy_formula_row(tmpl, src, write_row)
        else:
            ws.cell(row=write_row, column=column_index_from_string(COL_STOCK)).value = \
                f"=F{write_row}-K{write_row}"

        # N: 재고금액 — 수식 복사 또는 기본 =E*M
        if COL_STOCK_AMT in formula_tmpl:
            tmpl, src = formula_tmpl[COL_STOCK_AMT]
            ws.cell(row=write_row, column=column_index_from_string(COL_STOCK_AMT)).value = \
                _copy_formula_row(tmpl, src, write_row)
        else:
            ws.cell(row=write_row, column=column_index_from_string(COL_STOCK_AMT)).value = \
                f"=E{write_row}*M{write_row}"

        # P: 출처 (거래처명)
        if item.get("source"):
            src_cell = ws.cell(row=write_row, column=column_index_from_string(COL_SOURCE))
            src_cell.value     = item["source"]
            src_cell.font      = CELL_FONT
            src_cell.alignment = CELL_ALIGNMENT

        write_row += 1
        written   += 1

    wb.save(local_path)
    wb.close()
    logger.info("입고시트 %d행 작성 완료", written)
    return written


# ── 메인 ─────────────────────────────────────────────────────────────────────

def main() -> int:
    if not SUPABASE_URL or not SUPABASE_KEY:
        logger.error("SUPABASE_URL 또는 SUPABASE_SERVICE_ROLE_KEY 환경변수 없음")
        return 1

    if is_blackout():
        logger.info("매출 업데이트 시간대(00:05~00:20) — 스킵")
        return 0

    sessions = sb_get("inventory_sessions", {
        "status":        "eq.approved",
        "excel_updated": "eq.false",
        "select":        "*",
        "order":         "approved_at.asc",
    })

    if not sessions:
        logger.info("처리할 입고 세션 없음")
        return 0

    logger.info("처리 대기 입고 세션: %d건", len(sessions))
    dropbox  = DropboxClient.from_env("DROPBOX_ACCESS_TOKEN")
    WORK_DIR.mkdir(parents=True, exist_ok=True)

    for sess in sessions:
        session_id      = sess["id"]
        session_date    = sess["session_date"]
        barcode_prefix  = sess["barcode_prefix"]
        start_num_str   = str(sess.get("start_barcode_num") or "")
        location_def    = sess.get("location", "공동물류")
        is_session_online = (location_def == "온라인")

        # start_barcode_num은 레거시 fallback 용도 (일부 아이템에 barcode가 없을 때만 사용)
        has_start = start_num_str and start_num_str.isdigit()
        start_num = int(start_num_str) if has_start else 0
        num_width = len(start_num_str) if has_start else 6

        try:
            # 행거별 아이템 수집 (행거번호 → 순서대로)
            hangers = sb_get("inventory_hangers", {
                "session_id": f"eq.{session_id}",
                "select":     "*",
                "order":      "hanger_number.asc,created_at.asc",
            })
            hangers.sort(key=hanger_sort_key)

            ordered_items: list[dict] = []
            location = "온라인" if is_session_online else location_def
            for hanger in hangers:
                hanger_source = hanger.get("source") or ""   # 새 필드 (행거별 거래처)
                raw_items = sb_get("inventory_items", {
                    "hanger_id": f"eq.{hanger['id']}",
                    "select":    "*",
                    "order":     "order_index.asc",
                })

                for item in raw_items:
                    category_plain = item.get("category") or ""
                    apply_prefix   = is_session_online and category_plain not in NON_CLOTHING_CATEGORIES and barcode_prefix not in NO_ON_BARCODE_PREFIXES
                    category       = "온" + category_plain if apply_prefix else category_plain
                    if category_plain in ETC_MAPPED_CATEGORIES:
                        category = "ETC"

                    ordered_items.append({
                        "item_id":         item["id"],
                        "price":           item["price"],
                        "brand":           item.get("brand") or "",
                        "category":        category,        # H열용 (온라인이면 '온' 접두사 포함, 잡화 제외)
                        "category_plain":  category_plain,   # 상품명 조합용 ('온' 접두사 제외)
                        "location":        location,
                        "session_date":    session_date,
                        "existing_barcode": item.get("barcode") or "",   # 프론트에서 이미 부여됐으면 그대로 사용
                        "source":          hanger_source,               # P열용
                    })

            if not ordered_items:
                logger.info("아이템 없음, 완료 표시: %s", session_id[:8])
                sb_patch("inventory_sessions", {"id": session_id}, {
                    "excel_updated": True,
                    "status":        "processed",
                })
                continue

            # 바코드 부여: 프론트에서 이미 부여했으면 그대로, 아니면 레거시 방식 fallback
            missing_barcode_count = 0
            for i, item in enumerate(ordered_items):
                existing = item.get("existing_barcode") or ""
                if existing:
                    bc = existing
                else:
                    if not has_start:
                        raise ValueError(
                            f"바코드가 없는 아이템이 있는데 start_barcode_num도 없습니다: session={session_id}"
                        )
                    bc = f"{barcode_prefix}{str(start_num + i).zfill(num_width)}"
                    missing_barcode_count += 1
                brand = item["brand"]
                item["barcode"] = bc
                item["product_name"] = (
                    f"{bc} {brand} {item['category_plain']}" if brand
                    else f"{bc} {item['category_plain']}"
                )

            if missing_barcode_count:
                logger.info("레거시 바코드 부여: %d개 (session=%s)", missing_barcode_count, session_id[:8])

            # Supabase 아이템에 바코드/상품명 기록 (기존 값과 같아도 idempotent)
            for item in ordered_items:
                sb_patch("inventory_items", {"id": item["item_id"]}, {
                    "barcode":      item["barcode"],
                    "product_name": item["product_name"],
                })

            # Dropbox에서 엑셀 다운로드
            folder    = get_dropbox_folder(session_date)
            date_str  = session_date.replace("-", "")
            file_meta = dropbox.find_monthly_joint_sheet(folder, date_str)
            local_path = WORK_DIR / f"{session_id[:8]}_{file_meta.name}"
            dropbox.download_file(file_meta.path_display, local_path)

            # 입고 시트에 행 추가
            written = write_inventory_rows(local_path, ordered_items)

            # Dropbox 업로드 (rev 충돌 감지)
            if file_meta.rev:
                try:
                    dropbox.upload_with_rev_check(local_path, file_meta.path_display, file_meta.rev)
                except DropboxConflictError:
                    logger.warning("rev 충돌 — overwrite 모드로 재시도: %s", file_meta.path_display)
                    dropbox.upload_file(local_path, file_meta.path_display, mode="overwrite")
            else:
                dropbox.upload_file(local_path, file_meta.path_display, mode="overwrite")

            local_path.unlink(missing_ok=True)

            # 세션 완료 표시
            sb_patch("inventory_sessions", {"id": session_id}, {
                "excel_updated":    True,
                "excel_updated_at": datetime.now(timezone.utc).isoformat(),
                "total_items":      len(ordered_items),
                "status":           "processed",
            })

            first_bc = ordered_items[0]["barcode"]
            last_bc  = ordered_items[-1]["barcode"]
            logger.info(
                "✅ 입고 처리 완료: session=%s %d벌 (%s~%s)",
                session_id[:8], written, first_bc, last_bc,
            )

        except DropboxClientError as exc:
            logger.error("Dropbox 오류 (session=%s): %s", session_id[:8], exc)
        except Exception as exc:
            logger.exception("처리 실패 (session=%s): %s", session_id[:8], exc)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
